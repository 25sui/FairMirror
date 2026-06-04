from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".docx", ".pdf"}


def extract_document_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("仅支持 .txt、.md、.csv、.docx、.pdf 文件")

    if extension in {".txt", ".csv"}:
        return _decode_text(content)
    if extension == ".md":
        return _extract_markdown_text(content)
    if extension == ".docx":
        return _extract_docx_text(content)
    return _extract_pdf_text(content)


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = content.decode(encoding)
            return _normalize_text(text)
        except UnicodeDecodeError:
            continue
    return _normalize_text(content.decode("utf-8", errors="ignore"))


def _extract_markdown_text(content: bytes) -> str:
    text = _decode_text(content)
    text = re.sub(r"```[\s\S]*?```", "\n", text)
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"[*_`>~]", "", text)
    return _normalize_text(text)


def _extract_docx_text(content: bytes) -> str:
    try:
        text = _extract_docx_with_python_docx(content)
    except ImportError:
        text = ""
    except Exception:
        text = ""

    if text:
        return text

    return _extract_docx_from_xml(content)


def _extract_docx_with_python_docx(content: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(content))
    blocks: list[str] = []

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                blocks.append(paragraph.text)

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            if paragraph.text.strip():
                blocks.append(paragraph.text)

    return _normalize_text("\n".join(blocks))


def _extract_docx_from_xml(content: bytes) -> str:
    try:
        with zipfile.ZipFile(BytesIO(content)) as docx:
            xml_names = [
                "word/document.xml",
                "word/header1.xml",
                "word/header2.xml",
                "word/header3.xml",
                "word/footer1.xml",
                "word/footer2.xml",
                "word/footer3.xml",
            ]
            xml_parts = [docx.read(name) for name in xml_names if name in docx.namelist()]
    except (KeyError, zipfile.BadZipFile) as exc:
        raise ValueError("无法解析 docx 文件内容") from exc

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    blocks = []
    for xml in xml_parts:
        root = ElementTree.fromstring(xml)
        for paragraph in root.findall(".//w:p", namespace):
            text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace))
            if text.strip():
                blocks.append(text)

    text = _normalize_text("\n".join(blocks))
    if not text:
        raise ValueError("无法从 docx 文件中提取可审计文本")
    return text


def _extract_pdf_text(content: bytes) -> str:
    try:
        text = _extract_pdf_with_pdfplumber(content)
    except ImportError:
        text = ""
    except Exception:
        text = ""

    if text:
        return text

    return _extract_pdf_fallback(content)


def _extract_pdf_with_pdfplumber(content: bytes) -> str:
    import pdfplumber

    blocks: list[str] = []
    with pdfplumber.open(BytesIO(content)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            if page_text.strip():
                blocks.append(page_text)

            for table in page.extract_tables() or []:
                for row in table:
                    cells = [str(cell).strip() for cell in row if cell and str(cell).strip()]
                    if cells:
                        blocks.append(" | ".join(cells))

    return _normalize_text("\n".join(blocks))


def _extract_pdf_fallback(content: bytes) -> str:
    raw = content.decode("latin-1", errors="ignore")
    candidates = re.findall(r"\(([^()]{2,})\)", raw)
    text = "\n".join(item.replace("\\)", ")").replace("\\(", "(") for item in candidates)
    text = _normalize_text(text)
    if text:
        return text

    fallback = _normalize_text(re.sub(r"[^\x20-\x7E\u4e00-\u9fff]+", " ", raw))
    if len(fallback) < 20:
        raise ValueError("无法从 PDF 中提取可审计文本，建议上传可复制文本的 PDF；扫描件需要 OCR 能力")
    return fallback


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
